"""Alle Werte tab (moved verbatim from app.py main(), WP3e).

Renders the patient "Alle Werte" tab: the biomarker x report-date pivot matrix
with out-of-range / disabled highlighting, plus the xlsx export button.
"""

import io
from datetime import datetime

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Alignment
from openpyxl.utils import get_column_letter

from database import Report, DataPoint
from translations import t


def render_all_values_tab(session, patient, patient_reports):
    st.subheader(t("subheader_all_values"))
    st.caption(t("caption_all_values"))

    if patient_reports:
        # Gather all unique biomarkers (canonical_name + unit) for this patient.
        # Disabled measurements are included on purpose: they are shown in the
        # matrix with a yellow highlight so the user can see what is excluded.
        all_entries = (
            session.query(DataPoint.canonical_name, DataPoint.unit)
            .join(Report)
            .filter(
                Report.patient_id == patient.id,
                DataPoint.value.isnot(None),
            )
            .distinct()
            .order_by(DataPoint.canonical_name)
            .all()
        )

        # Gather all report dates for this patient (sorted ascending)
        report_dates = (
            session.query(Report.report_date)
            .filter(Report.patient_id == patient.id)
            .filter(Report.report_date.isnot(None))
            .distinct()
            .order_by(Report.report_date)
            .all()
        )
        date_list = [r[0] for r in report_dates]

        if all_entries and date_list:
            # Gather reference ranges (use the latest report's values per biomarker)
            ref_lookup = {}
            for cname, unit in all_entries:
                latest_ref = (
                    session.query(DataPoint.ref_low, DataPoint.ref_high)
                    .join(Report)
                    .filter(
                        Report.patient_id == patient.id,
                        DataPoint.canonical_name == cname,
                        DataPoint.unit == unit,
                        DataPoint.value.isnot(None),
                    )
                    .order_by(Report.report_date.desc())
                    .first()
                )
                ref_lookup[(cname, unit)] = (latest_ref.ref_low, latest_ref.ref_high) if latest_ref else (None, None)

            # Build a lookup dict: (canonical_name, unit, report_date) → value
            dp_query = (
                session.query(DataPoint.canonical_name, DataPoint.unit, DataPoint.value, Report.report_date, DataPoint.is_disabled)
                .join(Report)
                .filter(
                    Report.patient_id == patient.id,
                    DataPoint.value.isnot(None),
                )
                .all()
            )
            lookup = {}
            disabled_lookup = {}
            for cname, unit, value, rdate, is_dis in dp_query:
                lookup[(cname, unit, rdate)] = value
                disabled_lookup[(cname, unit, rdate)] = bool(is_dis)

            # Build pivot table rows with separate ref min/max columns
            pivot_rows = []
            disabled_by_label = {}
            for cname, unit in all_entries:
                label = f"{cname} ({unit})" if unit else cname
                ref_low, ref_high = ref_lookup.get((cname, unit), (None, None))
                row = {t("col_marker"): label, "Min": ref_low, "Max": ref_high}
                has_value = False
                for d in date_list:
                    val = lookup.get((cname, unit, d))
                    if val is not None:
                        row[d.strftime("%d.%m.%Y")] = val
                        has_value = True
                    else:
                        row[d.strftime("%d.%m.%Y")] = "—"
                # Only include rows that have at least one measured value
                if has_value:
                    pivot_rows.append(row)
                    disabled_by_label[label] = {
                        d.strftime("%d.%m.%Y") for d in date_list if disabled_lookup.get((cname, unit, d))
                    }

            df_pivot = pd.DataFrame(pivot_rows)

            # Display the pivot table with ref min/max columns
            display_df = df_pivot.set_index(t("col_marker")).reset_index()

            # Convert ALL columns to strings so pandas doesn't right-align numbers
            for col in display_df.columns:
                display_df[col] = display_df[col].apply(lambda x: str(x) if pd.notna(x) else '')

            def highlight_out_of_range(row, **kwargs):
                """Apply background color to cells where value is outside nominal range.

                Disabled measurements are highlighted yellow (they take
                precedence over the red/blue out-of-range colors).
                """
                from pandas import Series
                styles = Series(['text-align: center'] * len(row), index=row.index)

                label = row.get(t("col_marker"), '')
                disabled_dates = disabled_by_label.get(label, set())

                ref_low_str = row.get('Min', '')
                ref_high_str = row.get('Max', '')

                try:
                    ref_low = float(ref_low_str) if ref_low_str and ref_low_str != 'None' else None
                    ref_high = float(ref_high_str) if ref_high_str and ref_high_str != 'None' else None
                except (ValueError, TypeError):
                    return styles

                for col in row.index:
                    val = row[col]
                    if val in ('', '—') or pd.isna(val):
                        continue
                    if col in disabled_dates:
                        styles[col] = 'text-align: center; background-color: #f1c40f; color: #1a1a1a'
                        continue
                    if ref_low is None or ref_high is None:
                        continue
                    try:
                        num_val = float(val)
                        if num_val > ref_high:
                            styles[col] = 'text-align: center; background-color: #c0392b; color: #ffffff'
                        elif num_val < ref_low:
                            styles[col] = 'text-align: center; background-color: #2471a3; color: #ffffff'
                    except (ValueError, TypeError):
                        pass

                return styles

            styled_df = display_df.style.apply(highlight_out_of_range, axis=1)

            # Dynamic height: ~45px per row + header space, minimum 300px
            table_height = max(300, len(df_pivot) * 45 + 60)
            st.dataframe(styled_df, width='stretch', hide_index=True, height=table_height)

            # Export button
            col_export, _ = st.columns([1, 3])
            with col_export:
                if st.button(t("btn_export_csv"), type="secondary", width='stretch'):
                    wb = Workbook()
                    ws = wb.active
                    ws.title = t("sheet_all_values")[:31]  # Excel sheet-name limit

                    # Header row
                    headers = [t("col_marker"), t("col_reference_min"), t("col_reference_max")] + date_list
                    for col_idx, header in enumerate(headers, 1):
                        cell = ws.cell(row=1, column=col_idx, value=header)
                        cell.alignment = Alignment(horizontal="center")

                    # Data rows
                    dark_red_fill = PatternFill(start_color="C0392B", end_color="C0392B", fill_type="solid")
                    dark_blue_fill = PatternFill(start_color="2471A3", end_color="2471A3", fill_type="solid")
                    yellow_fill = PatternFill(start_color="F1C40F", end_color="F1C40F", fill_type="solid")
                    white_font_alignment = Alignment(horizontal="center", vertical="center")

                    for row_idx, row_data in enumerate(pivot_rows, 2):
                        marker_label = row_data[t("col_marker")]
                        ref_low = row_data.get('Min')
                        ref_high = row_data.get('Max')

                        # Marker column (left-aligned)
                        ws.cell(row=row_idx, column=1, value=marker_label).alignment = Alignment(horizontal="left", vertical="center")
                        # Reference columns
                        ws.cell(row=row_idx, column=2, value=ref_low if ref_low is not None else "").alignment = white_font_alignment
                        ws.cell(row=row_idx, column=3, value=ref_high if ref_high is not None else "").alignment = white_font_alignment

                        for col_idx, d in enumerate(date_list, 4):
                            date_str = d.strftime("%d.%m.%Y")
                            val = row_data.get(date_str, "—")
                            cell = ws.cell(row=row_idx, column=col_idx, value=val)
                            cell.alignment = white_font_alignment

                            if val == "—":
                                continue
                            if date_str in disabled_by_label.get(marker_label, set()):
                                cell.fill = yellow_fill
                                continue
                            try:
                                num_val = float(val)
                                if ref_low is not None and ref_high is not None:
                                    if num_val > float(ref_high):
                                        cell.fill = dark_red_fill
                                    elif num_val < float(ref_low):
                                        cell.fill = dark_blue_fill
                            except (ValueError, TypeError):
                                pass

                    # Column widths
                    ws.column_dimensions['A'].width = 30
                    for col_letter in ['B', 'C']:
                        ws.column_dimensions[col_letter].width = 14
                    for col_idx in range(4, len(headers) + 1):
                        ws.column_dimensions[get_column_letter(col_idx)].width = 14

                    # Save to bytes
                    xlsx_buffer = io.BytesIO()
                    wb.save(xlsx_buffer)
                    xlsx_bytes = xlsx_buffer.getvalue()

                    filename = f"{patient.name.replace(', ', '_')}_alle_werte_{datetime.now().strftime('%Y%m%d')}.xlsx"
                    st.download_button(
                        label=t("btn_export_csv"),
                        data=xlsx_bytes,
                        file_name=filename,
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        type="secondary",
                    )
                    st.success(t("msg_csv_exported", filename=filename))
        else:
            st.info(t("info_no_biomarker_data"))
    else:
        st.info(t("info_no_reports"))
